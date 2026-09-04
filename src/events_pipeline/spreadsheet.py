"""Reads and writes the events workbook.

Sheets:
  Events       - upcoming recommendations; you fill in the Feedback column.
  History      - events whose date has passed, moved out of Events to keep
                  the active list short. Feedback already learned from these
                  stays intact for the record.
  Preferences  - a read-only dump of the current learned weights, rewritten
                  every run so you can see what the model has picked up.

Event ID and Learned As are internal bookkeeping (hidden columns): Event ID
namespaces the source ("tm:" Ticketmaster, "ics:<feed>:" a public calendar
feed) so an id can never collide across sources, and is used to avoid ever
recommending the same event twice. Learned As records the Feedback value
already folded into the preference model, so a feedback change is only
learned from once. Feedback Source distinguishes feedback you typed
yourself from "Calendar (auto)" — inferred because a matching event showed
up on your own Google Calendar. Calendar Conflict flags events that
overlap something already on your calendar; it doesn't affect ranking.
"""
from __future__ import annotations

import datetime as dt
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .categorize import CategorizedEvent
from .preferences import Preferences

FEEDBACK_OPTIONS = ["Interested", "Maybe", "Not Interested", "Attended"]
AUTO_FEEDBACK_SOURCE = "Calendar (auto)"

EVENT_HEADERS = [
    "Event ID",
    "Date Added",
    "Event Name",
    "Category",
    "Event Date",
    "Venue",
    "City",
    "Genre",
    "Score",
    "Ticket URL",
    "Feedback",
    "Feedback Source",
    "Calendar Conflict",
    "Learned As",
]

(
    COL_EVENT_ID,
    COL_DATE_ADDED,
    COL_NAME,
    COL_CATEGORY,
    COL_EVENT_DATE,
    COL_VENUE,
    COL_CITY,
    COL_GENRE,
    COL_SCORE,
    COL_URL,
    COL_FEEDBACK,
    COL_FEEDBACK_SOURCE,
    COL_CONFLICT,
    COL_LEARNED_AS,
) = range(1, len(EVENT_HEADERS) + 1)

DATA_SHEETS = ("Events", "History")

PREF_HEADERS = ["Type", "Key", "Weight"]

HEADER_FILL = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def _format_data_sheet(ws) -> None:
    _style_header(ws, len(EVENT_HEADERS))
    ws.column_dimensions[get_column_letter(COL_EVENT_ID)].hidden = True
    ws.column_dimensions[get_column_letter(COL_LEARNED_AS)].hidden = True
    widths = {
        COL_DATE_ADDED: 12,
        COL_NAME: 42,
        COL_CATEGORY: 12,
        COL_EVENT_DATE: 20,
        COL_VENUE: 24,
        COL_CITY: 16,
        COL_GENRE: 14,
        COL_SCORE: 8,
        COL_URL: 40,
        COL_FEEDBACK: 15,
        COL_FEEDBACK_SOURCE: 16,
        COL_CONFLICT: 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _new_workbook() -> Workbook:
    wb = Workbook()
    events_ws = wb.active
    events_ws.title = "Events"
    events_ws.append(EVENT_HEADERS)
    _format_data_sheet(events_ws)

    history_ws = wb.create_sheet("History")
    history_ws.append(EVENT_HEADERS)
    _format_data_sheet(history_ws)

    prefs_ws = wb.create_sheet("Preferences")
    prefs_ws.append(PREF_HEADERS)
    _style_header(prefs_ws, len(PREF_HEADERS))

    _apply_feedback_validation(events_ws)
    return wb


def _apply_feedback_validation(ws) -> None:
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(FEEDBACK_OPTIONS)}"',
        allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(dv)
    col = get_column_letter(COL_FEEDBACK)
    # Generous range so the dropdown still applies to rows added later.
    dv.add(f"{col}2:{col}{max(ws.max_row, 2) + 2000}")


def open_workbook(path: str) -> Workbook:
    if os.path.exists(path):
        wb = load_workbook(path)
        if any(name not in wb.sheetnames for name in (*DATA_SHEETS, "Preferences")):
            raise ValueError(f"{path} exists but is missing expected sheets (Events/History/Preferences)")
        return wb
    return _new_workbook()


def save_workbook(wb: Workbook, path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def known_event_ids(wb: Workbook) -> set[str]:
    ids: set[str] = set()
    for sheet_name in DATA_SHEETS:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                ids.add(str(row[0]))
    return ids


def _parse_event_date(value) -> dt.date | None:
    if not value:
        return None
    try:
        return dt.datetime.fromisoformat(str(value).replace("Z", "+00:00")).date()
    except ValueError:
        return None


def pending_feedback_rows(wb: Workbook) -> list[dict]:
    """Rows (in Events or History) where Feedback differs from the
    last-learned value — covers both feedback you typed and feedback the
    attendance matcher auto-filled."""
    pending = []
    for sheet_name in DATA_SHEETS:
        ws = wb[sheet_name]
        for row_idx in range(2, ws.max_row + 1):
            event_id = ws.cell(row=row_idx, column=COL_EVENT_ID).value
            feedback = ws.cell(row=row_idx, column=COL_FEEDBACK).value
            learned_as = ws.cell(row=row_idx, column=COL_LEARNED_AS).value
            if not event_id or not feedback or feedback == learned_as:
                continue
            pending.append(
                {
                    "sheet": sheet_name,
                    "row": row_idx,
                    "event_id": str(event_id),
                    "name": ws.cell(row=row_idx, column=COL_NAME).value,
                    "category": ws.cell(row=row_idx, column=COL_CATEGORY).value,
                    "genre": ws.cell(row=row_idx, column=COL_GENRE).value,
                    "venue": ws.cell(row=row_idx, column=COL_VENUE).value,
                    "feedback": feedback,
                }
            )
    return pending


def mark_learned(wb: Workbook, sheet: str, row: int, feedback_value: str) -> None:
    wb[sheet].cell(row=row, column=COL_LEARNED_AS).value = feedback_value


def rows_needing_attendance_check(wb: Workbook) -> list[dict]:
    """Rows with no feedback yet, across both sheets — candidates the
    attendance matcher can propose "Attended" for."""
    candidates = []
    for sheet_name in DATA_SHEETS:
        ws = wb[sheet_name]
        for row_idx in range(2, ws.max_row + 1):
            event_id = ws.cell(row=row_idx, column=COL_EVENT_ID).value
            if not event_id:
                continue
            if ws.cell(row=row_idx, column=COL_FEEDBACK).value:
                continue
            candidates.append(
                {
                    "sheet": sheet_name,
                    "row": row_idx,
                    "name": ws.cell(row=row_idx, column=COL_NAME).value or "",
                    "event_date": _parse_event_date(ws.cell(row=row_idx, column=COL_EVENT_DATE).value),
                }
            )
    return candidates


def apply_attendance_matches(wb: Workbook, matches: list[dict]) -> None:
    for match in matches:
        ws = wb[match["sheet"]]
        ws.cell(row=match["row"], column=COL_FEEDBACK).value = "Attended"
        ws.cell(row=match["row"], column=COL_FEEDBACK_SOURCE).value = AUTO_FEEDBACK_SOURCE


def append_events(wb: Workbook, scored_events: list[tuple[CategorizedEvent, float, bool]], today: dt.date) -> None:
    ws = wb["Events"]
    for event, score, has_conflict in scored_events:
        ws.append(
            [
                event.event_id,
                today.isoformat(),
                event.name,
                event.category,
                event.start,
                event.venue,
                event.city,
                event.genre,
                score,
                event.url,
                None,
                None,
                "Yes" if has_conflict else None,
                None,
            ]
        )
    _apply_feedback_validation(ws)


def move_past_events_to_history(wb: Workbook, today: dt.date) -> int:
    events_ws = wb["Events"]
    history_ws = wb["History"]
    keep_rows = [EVENT_HEADERS]
    moved = 0
    for row in events_ws.iter_rows(min_row=2, values_only=True):
        event_date = _parse_event_date(row[COL_EVENT_DATE - 1])
        if event_date and event_date < today:
            history_ws.append(list(row))
            moved += 1
        else:
            keep_rows.append(list(row))

    events_ws.delete_rows(1, events_ws.max_row)
    for row in keep_rows:
        events_ws.append(row)
    _format_data_sheet(events_ws)
    _apply_feedback_validation(events_ws)
    return moved


def write_preferences_sheet(wb: Workbook, prefs: Preferences) -> None:
    ws = wb["Preferences"]
    ws.delete_rows(1, ws.max_row)
    ws.append(PREF_HEADERS)
    _style_header(ws, len(PREF_HEADERS))
    for row in prefs.to_rows():
        ws.append(row)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
